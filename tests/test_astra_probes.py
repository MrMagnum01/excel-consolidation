"""Astra's independent HOLD-review probes, wired into the repo's own test
suite so a regression is caught by `pytest tests`, not just by re-running a
one-off script by hand.

Source of truth for these cases: `~/vault/40-sessions/2026-09-27-astra-excel-consolidation-review.md`
and the accompanying `-probes.py` / `-probes.json`. Only filesystem paths are
adapted here (pytest's `tmp_path` instead of a manual `TemporaryDirectory`,
and a generated corpus instead of a hardcoded `data/workbooks` relative
path); the row/workbook construction and the behaviour each case demands are
exactly what the review specified. The assertions encode the review's
findings, not the pre-fix behaviour.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_consolidation.consolidator import process_directory
from excel_consolidation.generator import generate_all
from excel_consolidation.schema import CANONICAL_FIELDS, HEADER_STYLES

SEED = 7


def make(path: Path, rows: list[dict], extra: bool = False) -> None:
    """Build a one-sheet (or two-sheet, if `extra`) workbook with the given
    override rows on top of a synthetic base row. Mirrors probes.py's
    `make()` exactly."""
    w = Workbook()
    s = w.active
    s.title = "Orders"
    s.append([HEADER_STYLES[0][x] for x in CANONICAL_FIELDS])
    for row in rows:
        base = {
            "order_id": "o", "order_date": "2026-06-01", "customer": "Synthetic",
            "region": "Demo", "product": "Widget", "quantity": 2, "unit_price": 5,
            "amount": 10, "sales_rep": "Synthetic",
        }
        base.update(row)
        s.append([base[x] for x in CANONICAL_FIELDS])
        # Deliberately a literal string, NOT an input formula.
        if str(base["customer"]).startswith("="):
            s.cell(s.max_row, 3).data_type = "s"
    if extra:
        e = w.copy_worksheet(s)
        e.title = "More Orders"
        e.cell(2, 1).value = "second"
    w.save(path)


def _run_one(tmp_path: Path, name: str, rows: list[dict], extra: bool = False):
    inp = tmp_path / name
    inp.mkdir()
    make(inp / "input.xlsx", rows, extra)
    out = tmp_path / f"{name}-out"
    result = process_directory(inp, out)
    wb = load_workbook(out / "master.xlsx", data_only=False)
    data_ws = wb["Data"]
    header = [c.value for c in data_ws[1]]
    return result, data_ws, header


# --- Finding 1: non-finite money must never be accepted as OK ---------------

def test_nan_amount_is_rejected_not_accepted_as_blank(tmp_path):
    result, data_ws, _ = _run_one(tmp_path, "nan", [{"amount": "NaN"}])
    fr = result.file_results[0]
    assert fr.status != "ok"
    assert fr.rows_out == 0
    assert fr.rejected == 1
    assert {rr.reason for rr in result.rejected_rows} == {"invalid_amount"}
    assert data_ws.max_row == 1, "no data row should exist for a NaN amount"


def test_infinity_amount_is_rejected_not_accepted_as_blank(tmp_path):
    result, data_ws, _ = _run_one(tmp_path, "infinity", [{"amount": "Infinity"}])
    fr = result.file_results[0]
    assert fr.status != "ok"
    assert fr.rows_out == 0
    assert fr.rejected == 1
    assert {rr.reason for rr in result.rejected_rows} == {"invalid_amount"}
    assert data_ws.max_row == 1


def test_boolean_amount_is_rejected():
    from excel_consolidation.schema import parse_money
    assert parse_money(True) is None
    assert parse_money(False) is None


def test_negative_amount_is_out_of_domain_and_rejected():
    from excel_consolidation.schema import parse_money
    assert parse_money(-5) is None
    assert parse_money("-5.00") is None


# --- Finding 3: literal input text must never become an output formula -----

def test_formula_like_input_text_is_written_as_literal_output_text(tmp_path):
    result, data_ws, header = _run_one(tmp_path, "formula_text", [{"customer": "=1+1"}])
    assert result.file_results[0].rows_out == 1
    customer_col = header.index("Customer") + 1
    cell = data_ws.cell(2, customer_col)
    assert cell.data_type == "s", "customer text must be stored as a string, never a formula"
    assert cell.value == "=1+1", "the literal source text must survive unchanged"


# --- Finding 4: invalid amounts are rejected, never silently reconstructed -

def test_invalid_amount_is_rejected_not_silently_reconstructed_from_unit_price(tmp_path):
    result, data_ws, _ = _run_one(tmp_path, "bad_amount_replaced", [{"amount": "garbage"}])
    fr = result.file_results[0]
    assert fr.rows_out == 0
    assert fr.rejected == 1
    assert {rr.reason for rr in result.rejected_rows} == {"invalid_amount"}
    assert data_ws.max_row == 1, "no row should be reconstructed from unit_price x quantity"
    assert not any(k.startswith("amount_reconstructed") for k in result.change_log)


# --- Finding 2: currency identity must be preserved, not erased ------------

def test_explicit_currencies_are_preserved_as_a_currency_column(tmp_path):
    result, data_ws, header = _run_one(
        tmp_path, "currencies",
        [{"order_id": "a", "amount": "$10.00"}, {"order_id": "b", "amount": "€10.00"}],
    )
    assert result.file_results[0].rows_out == 2
    currency_col = header.index("Currency") + 1
    amount_col = header.index("Amount") + 1
    currencies = [data_ws.cell(r, currency_col).value for r in (2, 3)]
    amounts = [data_ws.cell(r, amount_col).value for r in (2, 3)]
    assert currencies == ["USD", "EUR"], "distinct currencies must not collapse into bare numbers"
    assert amounts == [10, 10]


def test_conflicting_amount_and_unit_price_currency_is_rejected(tmp_path):
    result, data_ws, _ = _run_one(
        tmp_path, "currency_mismatch",
        [{"amount": "$10.00", "unit_price": "€5.00"}],
    )
    fr = result.file_results[0]
    assert fr.rows_out == 0
    assert {rr.reason for rr in result.rejected_rows} == {"currency_mismatch"}
    assert data_ws.max_row == 1


def test_no_currency_symbol_mixing_within_a_single_generated_file(tmp_path):
    """Root-caused: the generator must pick one currency symbol per file,
    not per cell, or a single regional export ends up representing more
    than one currency in the same column -- exactly what the review found
    in 01_northgate.xlsx column 7 of the shipped corpus."""
    workbooks = tmp_path / "corpus"
    generate_all(workbooks, seed=SEED)
    mixed = []
    for p in sorted(workbooks.glob("*.xlsx")):
        try:
            wb = load_workbook(p, data_only=True)
        except Exception:
            continue
        for ws in wb:
            for col in ws.iter_cols():
                syms = {
                    c.value.strip()[0] for c in col
                    if isinstance(c.value, str) and c.value.strip() and c.value.strip()[0] in "$€£"
                }
                if len(syms) > 1:
                    mixed.append((p.name, col[0].column, sorted(syms)))
    assert mixed == [], f"generated corpus mixes currency symbols within a column: {mixed}"


# --- Finding 5: an extra data-shaped sheet must never disappear silently ---

def test_ambiguous_multi_sheet_workbook_is_excluded_not_silently_narrowed(tmp_path):
    result, data_ws, _ = _run_one(tmp_path, "extra_sheet", [{}], extra=True)
    fr = result.file_results[0]
    assert fr.status == "ambiguous_data_sheets"
    assert fr.rows_out == 0, "neither candidate sheet's rows should be silently accepted"
    assert {rr.reason for rr in result.rejected_rows} == {"ambiguous_data_sheets"}
    assert "files_ambiguous_data_sheets" in result.change_log
    assert data_ws.max_row == 1
