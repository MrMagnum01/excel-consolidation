"""End-to-end tests: generate the fixed-seed corpus, consolidate it, and
check the master workbook itself, idempotency, and the CLI."""

import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_consolidation.consolidator import process_directory
from excel_consolidation.generator import generate_all

SEED = 7
REPO_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="session")
def corpus_dir(tmp_path_factory):
    out = tmp_path_factory.mktemp("pl_workbooks")
    generate_all(out, seed=SEED)
    return out


def test_first_generated_order_lands_in_master(tmp_path_factory, corpus_dir):
    out = tmp_path_factory.mktemp("pl_out1")
    result = process_directory(corpus_dir, out)
    row = next(r for r in result.data_rows if r["order_id"] == "ORD-1001")
    assert row["source_file"] == "01_northgate.xlsx"
    assert row["region"] == "Northgate"


def test_master_workbook_has_the_three_expected_sheets(tmp_path_factory, corpus_dir):
    out = tmp_path_factory.mktemp("pl_out2")
    process_directory(corpus_dir, out)
    wb = load_workbook(out / "master.xlsx")
    assert wb.sheetnames == ["Data", "Validation", "Change log"]

    data_ws = wb["Data"]
    header = [c.value for c in next(data_ws.iter_rows(min_row=1, max_row=1))]
    assert header[:3] == ["Order Id", "Order Date", "Customer"]
    assert data_ws.max_row > 200  # a few hundred clean rows expected

    log_ws = wb["Change log"]
    log_header = [c.value for c in next(log_ws.iter_rows(min_row=1, max_row=1))]
    assert log_header == ["Transformation", "Count"]
    assert log_ws.max_row > 1


def test_region_inferred_from_filename_when_column_is_missing(tmp_path_factory, corpus_dir):
    out = tmp_path_factory.mktemp("pl_region")
    result = process_directory(corpus_dir, out)
    rows = [r for r in result.data_rows if r["source_file"] == "05_cedar_hollow.xlsx"]
    assert rows, "file 05 (region column intentionally omitted) produced no rows"
    assert all(r["region"] == "Cedar Hollow" for r in rows)


def test_idempotent_rerun_produces_identical_output(tmp_path_factory, corpus_dir):
    out1 = tmp_path_factory.mktemp("pl_idem1")
    out2 = tmp_path_factory.mktemp("pl_idem2")
    r1 = process_directory(corpus_dir, out1)
    r2 = process_directory(corpus_dir, out2)
    assert r1.data_rows == r2.data_rows
    assert r1.change_log == r2.change_log
    assert r1.totals == r2.totals

    # Re-running into the SAME output directory must not change the result.
    r3 = process_directory(corpus_dir, out1)
    assert r3.data_rows == r1.data_rows
    assert r3.totals == r1.totals


def test_cli_end_to_end_smoke(tmp_path):
    workbooks = tmp_path / "workbooks"
    out = tmp_path / "output"
    env = {**os.environ, "PYTHONPATH": str(REPO_ROOT / "src")}

    gen = subprocess.run(
        [sys.executable, "-m", "excel_consolidation", "generate", "--out", str(workbooks), "--seed", str(SEED)],
        cwd=str(REPO_ROOT), env=env, capture_output=True, text=True, timeout=60,
    )
    assert gen.returncode == 0, gen.stderr

    con = subprocess.run(
        [sys.executable, "-m", "excel_consolidation", "consolidate", "--input", str(workbooks), "--out", str(out)],
        cwd=str(REPO_ROOT), env=env, capture_output=True, text=True, timeout=60,
    )
    assert con.returncode == 0, con.stderr
    assert "Reconciled: True" in con.stdout
    assert (out / "master.xlsx").exists()
