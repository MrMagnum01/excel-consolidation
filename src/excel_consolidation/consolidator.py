"""Reads every workbook in a directory of messy regional order exports and
writes one clean master.xlsx with a Data sheet, a Validation sheet, and a
Change log sheet. Unreadable or wrong-structure files are reported as
categorised exceptions, never silently skipped, and the row counts always
reconcile: rows in = rows out + rejected + duplicates.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import ConsolidationResult, DuplicateRow, FileResult, RejectedRow
from .schema import (
    CURRENCY_SYMBOL_TO_CODE,
    DEFAULT_CURRENCY,
    REQUIRED_FIELDS,
    TOTAL_ROW_MARKERS,
    detect_currency_symbol,
    detect_field,
    is_blank_value,
    parse_date,
    parse_money,
    parse_quantity,
)

MIN_HEADER_MATCHES = 6


def _region_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    parts = stem.split("_")
    if parts and parts[0].isdigit():
        parts = parts[1:]
    label = " ".join(parts) if parts else stem
    return label.replace("-", " ").title()


def _find_header(ws) -> tuple[int | None, dict[int, str]]:
    """Scan the first few rows of a worksheet for the row with the most
    recognised canonical columns (robust to a merged banner/title row
    sitting above the real header). Returns (row_idx, {col: field}), or
    (None, {}) if no row reaches MIN_HEADER_MATCHES."""
    best_row, best_map = None, {}
    scan_limit = min(ws.max_row or 0, 6)
    for r in range(1, scan_limit + 1):
        mapping: dict[int, str] = {}
        for c in range(1, (ws.max_column or 0) + 1):
            fld = detect_field(ws.cell(row=r, column=c).value)
            if fld and fld not in mapping.values():
                mapping[c] = fld
        if len(mapping) > len(best_map):
            best_row, best_map = r, mapping
    if len(best_map) >= MIN_HEADER_MATCHES:
        return best_row, best_map
    return None, {}


def _pick_best_sheet(wb):
    """Pick the sheet with the most recognised header columns. If more than
    one sheet in the workbook independently reaches MIN_HEADER_MATCHES, the
    file is ambiguous -- which sheet is "the" order data can't be inferred
    silently, so the other viable sheets are returned too (each as its own
    (worksheet, header_row, col_map) candidate, not just a name) rather
    than picking one and dropping the rest unreported -- callers need the
    full candidate to count and reject that sheet's own rows too."""
    candidates: list[tuple] = []
    for ws in wb.worksheets:
        row, mapping = _find_header(ws)
        if row is not None:
            candidates.append((ws, row, mapping))
    if not candidates:
        return None, None, {}, []
    best_idx = max(range(len(candidates)), key=lambda i: len(candidates[i][2]))
    best_ws, best_row, best_map = candidates[best_idx]
    other_candidates = [c for i, c in enumerate(candidates) if i != best_idx]
    return best_ws, best_row, best_map, other_candidates


def process_directory(
    input_dir: Path, output_dir: Path, default_currency: str = DEFAULT_CURRENCY,
) -> ConsolidationResult:
    input_dir = Path(input_dir)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    data_rows: list[dict] = []
    file_results: list[FileResult] = []
    rejected_rows: list[RejectedRow] = []
    duplicate_rows: list[DuplicateRow] = []
    change_log: dict[str, int] = {}
    seen_order_ids: dict[str, tuple[str, int]] = {}

    def bump(key: str, n: int = 1) -> None:
        change_log[key] = change_log.get(key, 0) + n

    paths = sorted(
        p for p in input_dir.iterdir()
        if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")
    )

    for path in paths:
        filename = path.name
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable file is reported, never skipped
            file_results.append(FileResult(
                filename=filename, sheet_name=None, status="unreadable",
                rows_in=0, rows_out=0, rejected=0, duplicates=0,
                detail=f"{type(exc).__name__}: {exc}",
            ))
            bump("files_unreadable")
            continue

        ws, header_row, col_map, other_sheets = _pick_best_sheet(wb)
        if ws is None:
            probe_ws = wb.worksheets[0]
            n_rows = max((probe_ws.max_row or 1) - 1, 0)
            for r in range(2, (probe_ws.max_row or 0) + 1):
                rejected_rows.append(RejectedRow(
                    file=filename, row_ref=r, reason="unmappable_structure",
                    detail="no recognised order-data columns in this workbook",
                ))
            file_results.append(FileResult(
                filename=filename, sheet_name=probe_ws.title, status="unmappable_structure",
                rows_in=n_rows, rows_out=0, rejected=n_rows, duplicates=0,
                detail="headers did not match the order-data schema",
            ))
            bump("files_unmappable_structure")
            bump("rows_rejected_unmappable_structure", n_rows)
            continue

        if other_sheets:
            # More than one sheet in this workbook independently looks like
            # order data. Which one is authoritative can't be inferred, and
            # silently picking the "best" one means the other sheet's rows
            # vanish from every report. Reject the whole file instead of
            # guessing -- nothing here is silently dropped, it's an
            # explicit, named exception. Row coverage is counted across
            # *every* candidate sheet (not just the one that would have
            # been chosen): the file's total row count is the sum of each
            # candidate's own rows, and each rejected row records which
            # sheet it came from, since two candidates can each have a
            # "row 2" that are not the same row.
            other_titles = [c_ws.title for c_ws, _, _ in other_sheets]
            all_candidates = [(ws, header_row)] + [(c_ws, c_row) for c_ws, c_row, _ in other_sheets]
            n_rows = 0
            for c_ws, c_row in all_candidates:
                sheet_rows = max((c_ws.max_row or c_row) - c_row, 0)
                n_rows += sheet_rows
                for r in range(c_row + 1, (c_ws.max_row or c_row) + 1):
                    rejected_rows.append(RejectedRow(
                        file=filename, row_ref=r, reason="ambiguous_data_sheets",
                        detail=(
                            f"sheet {c_ws.title!r} (candidate sheets: "
                            f"{ws.title!r}, {', '.join(repr(s) for s in other_titles)})"
                        ),
                    ))
            file_results.append(FileResult(
                filename=filename, sheet_name=ws.title, status="ambiguous_data_sheets",
                rows_in=n_rows, rows_out=0, rejected=n_rows, duplicates=0,
                detail=(
                    f"{1 + len(other_sheets)} sheets look like order data "
                    f"(chose {ws.title!r}, also saw {', '.join(repr(s) for s in other_titles)}); "
                    "excluded rather than guessed; rows_in counts every candidate sheet"
                ),
            ))
            bump("files_ambiguous_data_sheets")
            bump("rows_rejected_ambiguous_data_sheets", n_rows)
            continue

        if header_row > 1:
            bump("banner_rows_skipped")
        if "region" not in col_map.values():
            bump("files_region_inferred_from_filename")
        fallback_region = _region_from_filename(filename)

        rows_in = rows_out = rejected = duplicates = 0
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            raw = {fld: ws.cell(row=r, column=c).value for c, fld in col_map.items()}
            rows_in += 1

            if all(is_blank_value(v) for v in raw.values()):
                rejected += 1
                rejected_rows.append(RejectedRow(file=filename, row_ref=r, reason="blank_row"))
                bump("blank_rows_removed")
                continue

            if any(isinstance(v, str) and v.strip().lower() in TOTAL_ROW_MARKERS for v in raw.values()):
                rejected += 1
                rejected_rows.append(RejectedRow(file=filename, row_ref=r, reason="totals_row"))
                bump("totals_rows_removed")
                continue

            missing = next((f for f in REQUIRED_FIELDS if is_blank_value(raw.get(f))), None)
            if missing:
                rejected += 1
                rejected_rows.append(RejectedRow(file=filename, row_ref=r, reason=f"missing_{missing}"))
                bump(f"rows_rejected_missing_{missing}")
                continue

            order_date = parse_date(raw["order_date"])
            if order_date is None:
                rejected += 1
                rejected_rows.append(RejectedRow(file=filename, row_ref=r, reason="invalid_date", detail=str(raw["order_date"])))
                bump("rows_rejected_invalid_date")
                continue
            if str(raw["order_date"]) != order_date.isoformat():
                bump("dates_normalized")

            quantity = parse_quantity(raw["quantity"])
            if quantity is None:
                rejected += 1
                rejected_rows.append(RejectedRow(file=filename, row_ref=r, reason="invalid_quantity", detail=str(raw["quantity"])))
                bump("rows_rejected_invalid_quantity")
                continue

            raw_amount = raw.get("amount")
            raw_unit_price = raw.get("unit_price")
            amount = parse_money(raw_amount)
            unit_price = parse_money(raw_unit_price)

            # `amount` is a REQUIRED_FIELDS entry, so a blank amount was
            # already rejected above as missing_amount. Reaching here with
            # amount is None means a value was *supplied* but didn't parse
            # (non-finite, negative, or not a number) -- that is rejected
            # outright. It is never silently reconstructed from unit_price:
            # a present-but-invalid amount is a data problem to report, not
            # a gap to paper over.
            if amount is None:
                rejected += 1
                rejected_rows.append(RejectedRow(file=filename, row_ref=r, reason="invalid_amount", detail=str(raw_amount)))
                bump("rows_rejected_invalid_amount")
                continue

            amount_symbol = detect_currency_symbol(raw_amount)
            price_symbol = detect_currency_symbol(raw_unit_price)
            if amount_symbol and price_symbol and amount_symbol != price_symbol:
                rejected += 1
                rejected_rows.append(RejectedRow(
                    file=filename, row_ref=r, reason="currency_mismatch",
                    detail=f"amount={amount_symbol!r} unit_price={price_symbol!r}",
                ))
                bump("rows_rejected_currency_mismatch")
                continue
            symbol = amount_symbol or price_symbol
            if symbol:
                currency = CURRENCY_SYMBOL_TO_CODE[symbol]
            else:
                currency = default_currency
                bump("currency_assumed_default")

            if unit_price is None:
                unit_price = round(amount / quantity, 2) if quantity else 0.0
                bump("unit_price_derived_from_amount_and_quantity")
            if isinstance(raw_amount, str) and any(sym in raw_amount for sym in "$€£¥"):
                bump("currency_symbols_stripped")
            if isinstance(raw_amount, str) and "," in raw_amount:
                bump("thousands_separators_removed")

            region = raw.get("region")
            if is_blank_value(region):
                region = fallback_region

            order_id = str(raw["order_id"]).strip()
            if order_id in seen_order_ids:
                first_file, first_row = seen_order_ids[order_id]
                duplicates += 1
                duplicate_rows.append(DuplicateRow(
                    file=filename, row_ref=r, order_id=order_id,
                    first_seen_file=first_file, first_seen_row=first_row,
                ))
                bump("duplicate_rows_removed")
                continue

            seen_order_ids[order_id] = (filename, r)
            data_rows.append({
                "order_id": order_id,
                "order_date": order_date.isoformat(),
                "customer": str(raw.get("customer") or "").strip(),
                "region": str(region).strip(),
                "product": str(raw.get("product") or "").strip(),
                "quantity": quantity,
                "unit_price": unit_price,
                "amount": amount,
                "currency": currency,
                "sales_rep": str(raw.get("sales_rep") or "").strip(),
                "source_file": filename,
                "source_row": r,
            })
            rows_out += 1

        status = "ok" if rejected == 0 and duplicates == 0 else "exception"
        file_results.append(FileResult(
            filename=filename, sheet_name=ws.title, status=status,
            rows_in=rows_in, rows_out=rows_out, rejected=rejected, duplicates=duplicates,
        ))

    result = ConsolidationResult(
        data_rows=data_rows, file_results=file_results,
        rejected_rows=rejected_rows, duplicate_rows=duplicate_rows,
        change_log=change_log,
    )
    _write_master_workbook(result, output_dir / "master.xlsx")
    return result


_DATA_COLUMNS = [
    "order_id", "order_date", "customer", "region", "product",
    "quantity", "unit_price", "amount", "currency", "sales_rep",
    "source_file", "source_row",
]


def _append_as_text(ws, values: list) -> None:
    """Append a row where every string value is forced to Excel's literal
    text type, never a formula. openpyxl infers a cell's type from its
    value when the value is assigned, so a source string like "=1+1"
    (typed by a customer into an input workbook, not an actual formula) is
    otherwise written out as an executable formula cell. Forcing
    `data_type = 's'` after assignment writes it as an inline/shared
    string instead, so it reloads (and displays in Excel) as the literal
    text it is."""
    ws.append(values)
    row_idx = ws.max_row
    for col_idx, value in enumerate(values, start=1):
        if isinstance(value, str):
            ws.cell(row=row_idx, column=col_idx).data_type = "s"


def _write_master_workbook(result: ConsolidationResult, path: Path) -> None:
    wb = Workbook()

    data_ws = wb.active
    data_ws.title = "Data"
    _append_as_text(data_ws, [c.replace("_", " ").title() for c in _DATA_COLUMNS])
    for row in result.data_rows:
        _append_as_text(data_ws, [row[c] for c in _DATA_COLUMNS])

    val_ws = wb.create_sheet("Validation")
    _append_as_text(val_ws, ["Per-file summary"])
    _append_as_text(val_ws, ["File", "Sheet", "Status", "Rows In", "Rows Out", "Rejected", "Duplicates", "Detail"])
    totals = result.totals
    for fr in result.file_results:
        _append_as_text(val_ws, [fr.filename, fr.sheet_name or "", fr.status, fr.rows_in, fr.rows_out, fr.rejected, fr.duplicates, fr.detail])
    val_ws.append([])
    _append_as_text(val_ws, [
        "TOTAL", "", "", totals["rows_in"], totals["rows_out"], totals["rejected"], totals["duplicates"],
        "reconciled" if result.reconciles() else "MISMATCH",
    ])
    val_ws.append([])
    _append_as_text(val_ws, ["Rejected rows"])
    _append_as_text(val_ws, ["File", "Row", "Reason", "Detail"])
    for rr in result.rejected_rows:
        _append_as_text(val_ws, [rr.file, rr.row_ref, rr.reason, rr.detail])
    val_ws.append([])
    _append_as_text(val_ws, ["Duplicates removed"])
    _append_as_text(val_ws, ["File", "Row", "Order ID", "First seen file", "First seen row"])
    for dr in result.duplicate_rows:
        _append_as_text(val_ws, [dr.file, dr.row_ref, dr.order_id, dr.first_seen_file, dr.first_seen_row])

    log_ws = wb.create_sheet("Change log")
    _append_as_text(log_ws, ["Transformation", "Count"])
    for key in sorted(result.change_log):
        _append_as_text(log_ws, [key, result.change_log[key]])

    wb.save(path)
