"""Deterministic generator for 30 messy synthetic regional order workbooks.

Builds:
  - 28 "regional" workbooks with realistic export mess: different header
    spellings/orders, a merged banner row above the header on some files,
    blank rows, an embedded totals row, dates as text in mixed formats,
    currency symbols, and duplicate records (within a file and across
    files).
  - 1 corrupt/unreadable file (not a valid xlsx at all).
  - 1 file with the wrong structure entirely (an inventory snapshot, not
    order data) that cannot be mapped onto the order schema.

Every row's *true* classification is tracked as it's built and returned as
a FileManifest per file, so tests can assert the consolidator's output
against known-correct ground truth rather than against itself.
"""

from __future__ import annotations

import datetime as _dt
import random
from pathlib import Path

from openpyxl import Workbook

from .models import FileManifest
from .regions import FIRST_NAMES, INVENTORY_LOCATIONS, LAST_NAMES, PRODUCTS, REGIONS
from .schema import DATE_FORMATS, HEADER_STYLE_ORDERS, HEADER_STYLES

NUM_FILES = 30
CORRUPT_POSITION = 15
WRONG_STRUCTURE_POSITION = 27

INVALID_DATE_STRINGS = ["31/02/2026", "not a date", "00/00/0000"]
INVALID_QUANTITY_STRINGS = ["many", "n/a", "-"]
INVALID_AMOUNT_STRINGS = ["TBD", "call for price", "???"]

ALL_DEFECT_TYPES = [
    "missing_order_id", "missing_order_date", "missing_product",
    "missing_quantity", "missing_amount", "invalid_date",
    "invalid_quantity", "invalid_amount",
]

# Guarantees every rejection category appears at least once in the corpus.
FORCED_DEFECTS = {
    1: "missing_order_id",
    2: "missing_order_date",
    3: "missing_product",
    4: "missing_quantity",
    5: "missing_amount",
    6: "invalid_date",
    7: "invalid_quantity",
    8: "invalid_amount",
}


def _fmt_money(value: float, symbol: str, rng: random.Random) -> str:
    if rng.random() < 0.4:
        text = f"{value:,.2f}"
    else:
        text = f"{value:.2f}"
    return f"{symbol}{text}"


def _fmt_date(d: _dt.date, rng: random.Random) -> str:
    return d.strftime(rng.choice(DATE_FORMATS))


def _random_date(rng: random.Random) -> _dt.date:
    start = _dt.date(2026, 1, 1)
    d = start + _dt.timedelta(days=rng.randint(0, 179))
    if d.day > 28:
        d = d.replace(day=28)
    return d


def _person(rng: random.Random) -> str:
    return f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)}"


def _order_id(counter: list[int]) -> str:
    counter[0] += 1
    return f"ORD-{1000 + counter[0]}"


def _base_row(region: str, rng: random.Random, id_counter: list[int]) -> dict:
    qty = rng.randint(1, 40)
    price = round(rng.uniform(5, 500), 2)
    return {
        "order_id": _order_id(id_counter),
        "order_date": _random_date(rng),
        "customer": _person(rng),
        "region": region,
        "product": rng.choice(PRODUCTS),
        "quantity": qty,
        "unit_price": price,
        "amount": round(qty * price, 2),
        "sales_rep": _person(rng),
    }


def _make_invalid_row(defect: str, region: str, rng: random.Random, id_counter: list[int]) -> dict:
    row = _base_row(region, rng, id_counter)
    if defect == "missing_order_id":
        row["order_id"] = None
    elif defect == "missing_order_date":
        row["order_date"] = None
    elif defect == "missing_product":
        row["product"] = None
    elif defect == "missing_quantity":
        row["quantity"] = None
        row["amount"] = None
    elif defect == "missing_amount":
        row["amount"] = None
        row["unit_price"] = None
    elif defect == "invalid_date":
        row["order_date"] = rng.choice(INVALID_DATE_STRINGS)
    elif defect == "invalid_quantity":
        row["quantity"] = rng.choice(INVALID_QUANTITY_STRINGS)
    elif defect == "invalid_amount":
        row["amount"] = rng.choice(INVALID_AMOUNT_STRINGS)
        row["unit_price"] = None
    return row


def _cell_value(field: str, value: object, currency_symbol: str, rng: random.Random):
    if value is None:
        return None
    if field == "order_date":
        return _fmt_date(value, rng) if isinstance(value, _dt.date) else value
    if field in ("unit_price", "amount"):
        return _fmt_money(value, currency_symbol, rng) if isinstance(value, (int, float)) else value
    if field == "quantity" and isinstance(value, int):
        return str(value) if rng.random() < 0.3 else value
    return value


def _build_normal_file_rows(
    slot: int, region: str, rng: random.Random, id_counter: list[int],
    all_emitted: list[dict], forced_defect: str | None,
) -> tuple[list[dict], FileManifest, int, bool]:
    style_idx = slot % len(HEADER_STYLES)
    has_region_column = slot % 5 != 0

    n_valid = rng.randint(10, 18)
    valid_rows = [_base_row(region, rng, id_counter) for _ in range(n_valid)]
    events = [{"kind": "valid", "data": r} for r in valid_rows]

    duplicates_added = 0
    if slot > 1 and all_emitted and rng.random() < 0.5:
        events.append({"kind": "duplicate", "data": dict(rng.choice(all_emitted))})
        duplicates_added += 1
    if valid_rows and rng.random() < 0.3:
        events.append({"kind": "duplicate", "data": dict(rng.choice(valid_rows))})
        duplicates_added += 1

    rejected_counts: dict[str, int] = {}

    n_blank = rng.choice([0, 0, 1, 1, 2])
    for _ in range(n_blank):
        events.append({"kind": "blank", "data": None})
    if n_blank:
        rejected_counts["blank_row"] = n_blank

    if rng.random() < 0.6:
        total_amount = round(sum(r["amount"] for r in valid_rows), 2)
        marker = rng.choice(["TOTAL", "Grand Total", "Subtotal", "TOTALS"])
        events.append({"kind": "totals", "data": {"customer": marker, "amount": total_amount}})
        rejected_counts["totals_row"] = 1

    defects = [forced_defect] if forced_defect else []
    defects.extend(rng.choice(ALL_DEFECT_TYPES) for _ in range(rng.choice([0, 0, 1])))
    for defect in defects:
        row = _make_invalid_row(defect, region, rng, id_counter)
        events.append({"kind": "invalid", "data": row})
        rejected_counts[defect] = rejected_counts.get(defect, 0) + 1

    rng.shuffle(events)
    # openpyxl silently drops a wholly-empty *trailing* row on save/reload,
    # which would undercount rows_in vs. the manifest below -- keep a
    # non-blank event last so every intentional blank row round-trips.
    if events and events[-1]["kind"] == "blank":
        for i in range(len(events) - 1):
            if events[i]["kind"] != "blank":
                events[i], events[-1] = events[-1], events[i]
                break

    all_emitted.extend(valid_rows)

    manifest = FileManifest(
        filename="", kind="normal", rows_in=len(events),
        valid_unique=len(valid_rows), duplicates=duplicates_added,
        rejected=rejected_counts,
    )
    return events, manifest, style_idx, has_region_column


def _write_normal_workbook(
    path: Path, sheet_name: str, header_style: dict, field_order: list[str],
    add_banner: bool, events: list[dict], render_rng: random.Random,
    currency_symbol: str,
) -> None:
    """`render_rng` is a stream isolated from the content-decision `rng`
    used by _build_normal_file_rows(): it only chooses *cosmetic* things
    (which date format to render, which thousands-separator style) and
    must never be the same object as the content rng, or build_manifests()
    (which never renders a cell) would silently drift out of sync with
    generate_all() (which does).

    `currency_symbol` is fixed for the whole file (chosen once by the
    caller), not per cell: a real regional export uses one currency
    throughout, and this workbook's `amount`/`unit_price` columns must not
    mix symbols within themselves, or the file no longer represents a
    single-currency source and the consolidator has no sound basis for
    picking one."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ncols = len(field_order)

    row_idx = 1
    if add_banner:
        ws.cell(row=1, column=1, value="Sales Export Summary")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        row_idx = 2

    header_row = row_idx
    for col, fld in enumerate(field_order, start=1):
        ws.cell(row=header_row, column=col, value=header_style[fld])
    row_idx += 1

    for event in events:
        data = event["data"]
        for col, fld in enumerate(field_order, start=1):
            value = None if data is None else _cell_value(fld, data.get(fld), currency_symbol, render_rng)
            ws.cell(row=row_idx, column=col, value=value)
        row_idx += 1

    wb.save(path)


def _wrong_structure_rows(rng: random.Random) -> list[list]:
    """The RNG-consuming part of the inventory-snapshot file, factored out
    so build_manifests() (no files written) stays in lockstep with
    generate_all() (files written) on every later call to `rng`."""
    rows = []
    for i in range(10):
        d = _random_date(rng)
        rows.append([
            f"SKU-{2000 + i}",
            rng.choice(INVENTORY_LOCATIONS),
            rng.randint(0, 500),
            rng.randint(10, 50),
            d.strftime("%Y-%m-%d"),
        ])
    return rows


def _write_wrong_structure_workbook(path: Path, rows: list[list]) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    headers = ["SKU", "Warehouse", "Stock Level", "Reorder Point", "Last Restocked"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)
    for i, row in enumerate(rows):
        for col, v in enumerate(row, start=1):
            ws.cell(row=i + 2, column=col, value=v)
    wb.save(path)
    return len(rows)


def generate_all(out_dir: Path, seed: int = 7) -> list[FileManifest]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    render_rng = random.Random(f"render:{seed}")  # cosmetic-only; see _write_normal_workbook
    id_counter = [0]
    all_emitted: list[dict] = []
    manifests: list[FileManifest] = []
    region_iter = iter(REGIONS)

    for slot in range(1, NUM_FILES + 1):
        if slot == CORRUPT_POSITION:
            filename = f"{slot:02d}_corrupt_export.xlsx"
            (out_dir / filename).write_bytes(
                b"NOT A VALID XLSX FILE - simulated corrupted regional export\x00\x01\x02"
            )
            manifests.append(FileManifest(
                filename=filename, kind="corrupt", rows_in=0,
                valid_unique=0, duplicates=0, rejected={},
            ))
            continue

        if slot == WRONG_STRUCTURE_POSITION:
            filename = f"{slot:02d}_inventory_snapshot.xlsx"
            rows = _wrong_structure_rows(rng)
            n_rows = _write_wrong_structure_workbook(out_dir / filename, rows)
            manifests.append(FileManifest(
                filename=filename, kind="wrong_structure", rows_in=n_rows,
                valid_unique=0, duplicates=0,
                rejected={"unmappable_structure": n_rows},
            ))
            continue

        region = next(region_iter)
        filename = f"{slot:02d}_{region.lower().replace(' ', '_')}.xlsx"
        forced_defect = FORCED_DEFECTS.get(slot)
        events, manifest, style_idx, has_region_column = _build_normal_file_rows(
            slot, region, rng, id_counter, all_emitted, forced_defect,
        )
        manifest.filename = filename
        field_order = [f for f in HEADER_STYLE_ORDERS[style_idx] if has_region_column or f != "region"]
        header_style = HEADER_STYLES[style_idx]
        add_banner = slot % 3 == 0
        sheet_name = render_rng.choice(["Data", "Sales", "Sheet1", "Export", region[:20]])
        # One currency symbol for the whole file -- see _write_normal_workbook.
        currency_symbol = render_rng.choice(["", "", "$", "€", "£"])
        _write_normal_workbook(
            out_dir / filename, sheet_name, header_style, field_order,
            add_banner, events, render_rng, currency_symbol,
        )
        manifests.append(manifest)

    return manifests


def build_manifests(seed: int = 7) -> list[FileManifest]:
    """Ground truth only, without writing any files (used by fast tests)."""
    rng = random.Random(seed)
    id_counter = [0]
    all_emitted: list[dict] = []
    manifests: list[FileManifest] = []
    region_iter = iter(REGIONS)

    for slot in range(1, NUM_FILES + 1):
        if slot == CORRUPT_POSITION:
            manifests.append(FileManifest(
                filename=f"{slot:02d}_corrupt_export.xlsx", kind="corrupt",
                rows_in=0, valid_unique=0, duplicates=0, rejected={},
            ))
            continue
        if slot == WRONG_STRUCTURE_POSITION:
            rows = _wrong_structure_rows(rng)
            manifests.append(FileManifest(
                filename=f"{slot:02d}_inventory_snapshot.xlsx", kind="wrong_structure",
                rows_in=len(rows), valid_unique=0, duplicates=0,
                rejected={"unmappable_structure": len(rows)},
            ))
            continue
        region = next(region_iter)
        filename = f"{slot:02d}_{region.lower().replace(' ', '_')}.xlsx"
        forced_defect = FORCED_DEFECTS.get(slot)
        _events, manifest, _style_idx, _has_region = _build_normal_file_rows(
            slot, region, rng, id_counter, all_emitted, forced_defect,
        )
        manifest.filename = filename
        manifests.append(manifest)

    return manifests
